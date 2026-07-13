"""Unit tests for Excel ↔ UMF bidirectional converter."""

from __future__ import annotations

import json

import openpyxl
import pytest

from tablespec.excel_converter import (
    ExcelConstants,
    ExcelToUMFConverter,
    ExcelValidator,
    UMFToExcelConverter,
)
from tablespec.models.umf import (
    UMF,
    DerivationCandidate,
    ForeignKey,
    FileFormatSpec,
    FilenamePattern,
    JdbcSource,
    JoinViaSpec,
    Nullable,
    OutgoingRelationship,
    Relationships,
    Survivorship,
    UMFColumn,
    UMFColumnDerivation,
    ValidationRules,
)

pytestmark = [pytest.mark.no_spark, pytest.mark.fast]


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _make_minimal_umf(**overrides) -> UMF:
    """Create a minimal valid UMF for testing."""
    defaults = {
        "version": "1.0",
        "table_name": "test_table",
        "canonical_name": "Test Table",
        "columns": [
            UMFColumn(
                name="col_id",
                data_type="INTEGER",
                description="Primary key",
            ),
            UMFColumn(
                name="col_name",
                data_type="VARCHAR",
                length=100,
                description="Name field",
                nullable=Nullable(MD=True, MP=False, ME=True),
                sample_values=["Alice", "Bob"],
            ),
        ],
    }
    defaults.update(overrides)
    return UMF(**defaults)


def _make_rich_umf(**overrides) -> UMF:
    """Create a UMF with validation rules, relationships, and various column types."""
    defaults = {
        "version": "2.0",
        "table_name": "rich_table",
        "canonical_name": "Rich Table",
        "description": "A table with many features",
        "table_type": "provided",
        "primary_key": ["col_id"],
        "aliases": ["rich", "Rich Table"],
        "columns": [
            UMFColumn(
                name="col_id",
                data_type="INTEGER",
                description="Primary key",
                key_type="primary",
                source="data",
            ),
            UMFColumn(
                name="col_name",
                data_type="VARCHAR",
                length=200,
                description="Full name",
                nullable=Nullable(MD=True, MP=True, ME=False),
                sample_values=["Alice", "Bob"],
                canonical_name="Full Name",
                aliases=["name", "full_name"],
                reporting_requirement="R",
                format="Last, First",
            ),
            UMFColumn(
                name="col_amount",
                data_type="DECIMAL",
                precision=10,
                scale=2,
                description="Dollar amount",
            ),
            UMFColumn(
                name="col_active",
                data_type="BOOLEAN",
                description="Is active flag",
            ),
            UMFColumn(
                name="col_date",
                data_type="DATE",
                description="Created date",
                format="YYYY-MM-DD",
                notes=["Must be after 2020-01-01", "Business calendar only"],
            ),
        ],
        "validation_rules": ValidationRules(
            expectations=[
                {
                    "type": "expect_column_values_to_not_be_null",
                    "kwargs": {"column": "col_id"},
                    "meta": {
                        "description": "ID must not be null",
                        "severity": "critical",
                    },
                },
                {
                    "type": "expect_column_values_to_be_in_set",
                    "kwargs": {"column": "col_active", "value_set": [True, False]},
                    "meta": {
                        "description": "Must be boolean",
                        "severity": "warning",
                    },
                },
            ],
        ),
        "relationships": Relationships(
            foreign_keys=[
                ForeignKey(
                    column="col_id",
                    references_table="other_table",
                    references_column="id",
                ),
            ],
            outgoing=[
                OutgoingRelationship(
                    source_column="col_id",
                    target_table="other_table",
                    target_column="id",
                    type="foreign_to_primary",
                    confidence=0.95,
                    reasoning="ID match",
                ),
            ],
        ),
    }
    defaults.update(overrides)
    return UMF(**defaults)


def _make_manual_workbook(column_rows: list[dict[str, object]]) -> openpyxl.Workbook:
    """Create a minimal Excel workbook for importer-focused tests."""
    wb = openpyxl.Workbook()
    schema_ws = wb.active
    schema_ws.title = ExcelConstants.SHEET_SCHEMA
    schema_ws["A1"] = "Field"
    schema_ws["B1"] = "Value"
    schema_ws["A2"] = "table_name"
    schema_ws["B2"] = "alias_table"
    schema_ws["A3"] = "canonical_name"
    schema_ws["B3"] = "Alias Table"

    columns_ws = wb.create_sheet(ExcelConstants.SHEET_COLUMNS)
    headers = [
        "Name",
        "Canonical Name",
        "Aliases",
        "Data Type",
        "Length",
        "Precision",
        "Scale",
    ]
    for idx, header in enumerate(headers, start=1):
        columns_ws.cell(1, idx).value = header

    for row_idx, row in enumerate(column_rows, start=2):
        columns_ws.cell(row_idx, 1).value = row["name"]
        columns_ws.cell(row_idx, 2).value = row.get("canonical_name")
        columns_ws.cell(row_idx, 3).value = row.get("aliases")
        columns_ws.cell(row_idx, 4).value = row.get("data_type")
        columns_ws.cell(row_idx, 5).value = row.get("length")
        columns_ws.cell(row_idx, 6).value = row.get("precision")
        columns_ws.cell(row_idx, 7).value = row.get("scale")

    return wb


@pytest.fixture()
def minimal_umf():
    return _make_minimal_umf()


@pytest.fixture()
def rich_umf():
    return _make_rich_umf()


# ---------------------------------------------------------------------------
# UMFToExcelConverter tests
# ---------------------------------------------------------------------------


class TestUMFToExcelConverter:
    """Test UMF -> Excel export."""

    def test_convert_returns_workbook(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        assert isinstance(wb, openpyxl.Workbook)

    def test_required_sheets_present(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        names = wb.sheetnames
        assert ExcelConstants.SHEET_README in names
        assert ExcelConstants.SHEET_SCHEMA in names
        assert ExcelConstants.SHEET_COLUMNS in names
        assert ExcelConstants.SHEET_METADATA in names

    def test_schema_sheet_values(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        ws = wb[ExcelConstants.SHEET_SCHEMA]
        # Row 2 = table_name, Row 3 = canonical_name
        assert ws["A2"].value == "table_name"
        assert ws["B2"].value == "test_table"
        assert ws["A3"].value == "canonical_name"
        assert ws["B3"].value == "Test Table"

    def test_columns_sheet_headers(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # Check header row
        assert ws["A1"].value == "Name"
        assert ws["B1"].value == "Canonical Name"
        assert ws["D1"].value == "Data Type"

    def test_columns_sheet_data(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # Row 2 = first column
        assert ws["A2"].value == "col_id"
        assert ws["D2"].value == "INTEGER"
        # Row 3 = second column
        assert ws["A3"].value == "col_name"
        assert ws["D3"].value == "VARCHAR"
        assert ws["E3"].value == 100  # length

    def test_nullable_values_in_columns_sheet(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # col_name (row 3) has nullable MD=True, MP=False, ME=True
        # Context keys are sorted alphabetically: MD, ME, MP
        assert ws["H3"].value is True  # MD
        assert ws["I3"].value is True  # ME
        assert ws["J3"].value is False  # MP

    def test_sample_values_in_columns_sheet(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # col_name has sample_values ["Alice", "Bob"]
        assert ws["L3"].value == "Alice, Bob"

    def test_validation_sheet_created_when_rules_exist(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        assert ExcelConstants.SHEET_VALIDATION in wb.sheetnames

    def test_validation_sheet_not_created_when_no_rules(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        assert ExcelConstants.SHEET_VALIDATION not in wb.sheetnames

    def test_relationships_sheet_created_when_fks_exist(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        assert ExcelConstants.SHEET_RELATIONSHIPS in wb.sheetnames

    def test_relationships_sheet_not_created_when_no_fks(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        assert ExcelConstants.SHEET_RELATIONSHIPS not in wb.sheetnames

    def test_save_to_file(self, minimal_umf, tmp_path):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        out = tmp_path / "test.xlsx"
        wb.save(out)
        assert out.exists()
        # Re-open to ensure file is valid
        wb2 = openpyxl.load_workbook(out)
        assert ExcelConstants.SHEET_SCHEMA in wb2.sheetnames

    def test_description_in_schema(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_SCHEMA]
        # Find description row
        found = False
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value == "description":
                assert row[1].value == "A table with many features"
                found = True
                break
        assert found, "description field not found in Schema sheet"

    def test_decimal_column_precision_scale(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # Find col_amount row
        for row_cells in ws.iter_rows(min_row=2, values_only=False):
            if row_cells[0].value == "col_amount":
                assert row_cells[3].value == "DECIMAL"  # data_type
                assert row_cells[4].value == ""  # length (not applicable)
                assert row_cells[5].value == 10  # precision
                assert row_cells[6].value == 2  # scale
                return
        pytest.fail("col_amount not found in Columns sheet")

    def test_notes_exported_as_newline_string(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # Find col_date row
        for row_cells in ws.iter_rows(min_row=2, values_only=False):
            if row_cells[0].value == "col_date":
                notes_val = row_cells[17].value  # R column = index 17
                assert "Must be after 2020-01-01" in notes_val
                assert "Business calendar only" in notes_val
                return
        pytest.fail("col_date not found in Columns sheet")

    def test_format_exported(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        for row_cells in ws.iter_rows(min_row=2, values_only=False):
            if row_cells[0].value == "col_name":
                assert row_cells[16].value == "Last, First"  # Q column = index 16
                return
        pytest.fail("col_name not found in Columns sheet")

    def test_primary_key_in_schema(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_SCHEMA]
        found = False
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value == "primary_key":
                assert row[1].value == "col_id"
                found = True
                break
        assert found, "primary_key not found in Schema sheet"

    def test_aliases_in_columns(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        for row_cells in ws.iter_rows(min_row=2, values_only=False):
            if row_cells[0].value == "col_name":
                aliases_val = row_cells[2].value  # C column = aliases
                assert "name" in aliases_val
                assert "full_name" in aliases_val
                return
        pytest.fail("col_name not found")


# ---------------------------------------------------------------------------
# ExcelValidator tests
# ---------------------------------------------------------------------------


class TestExcelValidator:
    """Test Excel workbook validation."""

    def test_valid_workbook_passes(self, minimal_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(minimal_umf)
        validator = ExcelValidator()
        assert validator.validate_workbook(wb) is True

    def test_missing_schema_sheet_fails(self):
        wb = openpyxl.Workbook()
        # Only default sheet, no Schema or Columns
        validator = ExcelValidator()
        assert validator.validate_workbook(wb) is False
        assert any("Schema" in e for e in validator.errors)

    def test_missing_columns_sheet_fails(self):
        wb = openpyxl.Workbook()
        wb.create_sheet("Schema")
        ws = wb["Schema"]
        ws["A2"] = "table_name"
        ws["B2"] = "test"
        ws["A3"] = "canonical_name"
        ws["B3"] = "Test"
        validator = ExcelValidator()
        assert validator.validate_workbook(wb) is False
        assert any("Columns" in e for e in validator.errors)

    def test_rich_workbook_passes(self, rich_umf):
        converter = UMFToExcelConverter()
        wb = converter.convert(rich_umf)
        validator = ExcelValidator()
        result = validator.validate_workbook(wb)
        # If there are errors, print them for debugging
        if not result:
            pytest.fail(f"Validation failed: {validator.errors}")


# ---------------------------------------------------------------------------
# ExcelToUMFConverter tests (individual extraction methods)
# ---------------------------------------------------------------------------


class TestExcelToUMFConverterExtractSchema:
    """Test _extract_schema method."""

    def test_extract_schema_from_exported_workbook(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        schema = importer._extract_schema(wb)
        assert schema["table_name"] == "rich_table"
        assert schema["canonical_name"] == "Rich Table"
        assert schema["description"] == "A table with many features"
        assert schema["version"] == "2.0"
        assert schema.get("table_type") == "provided"

    def test_extract_schema_aliases(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        schema = importer._extract_schema(wb)
        # Aliases should be parsed as list; "Rich Table" filtered out by exporter since it matches canonical_name
        if "aliases" in schema:
            assert isinstance(schema["aliases"], list)

    def test_extract_schema_primary_key(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        schema = importer._extract_schema(wb)
        assert "primary_key" in schema
        assert "col_id" in schema["primary_key"]


class TestExcelToUMFConverterExtractColumns:
    """Test _extract_columns method."""

    def test_extract_columns_from_exported_workbook(self, minimal_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert len(columns) == 2
        assert columns[0]["name"] == "col_id"
        assert columns[1]["name"] == "col_name"

    def test_extract_columns_data_types(self, minimal_umf):
        """Data types are normalized to UMF spellings by the importer."""
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert columns[0]["data_type"] == "INTEGER"
        assert columns[1]["data_type"] == "VARCHAR"

    def test_extract_columns_normalizes_type_aliases(self):
        wb = _make_manual_workbook(
            [
                {"name": "col_int_raw", "data_type": "INTEGER"},
                {"name": "col_int_alias", "data_type": "IntegerType"},
                {"name": "col_str_raw", "data_type": "VARCHAR", "length": 50},
                {"name": "col_str_alias", "data_type": "StringType", "length": 50},
                {
                    "name": "col_dec_raw",
                    "data_type": "DECIMAL",
                    "precision": 10,
                    "scale": 2,
                },
                {
                    "name": "col_dec_alias",
                    "data_type": "DecimalType",
                    "precision": 12,
                    "scale": 4,
                },
            ]
        )
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert [col["data_type"] for col in columns] == [
            "INTEGER",
            "INTEGER",
            "VARCHAR",
            "VARCHAR",
            "DECIMAL",
            "DECIMAL",
        ]

    def test_extract_columns_length(self, minimal_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert columns[1].get("length") == 100

    def test_extract_columns_nullable(self, minimal_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        # col_name has nullable MD=True, MP=False, ME=True
        nullable = columns[1].get("nullable")
        assert nullable is not None
        assert nullable["MD"] is True
        assert nullable["MP"] is False
        assert nullable["ME"] is True

    def test_extract_columns_sample_values(self, minimal_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        sv = columns[1].get("sample_values")
        assert sv is not None
        assert "Alice" in sv
        assert "Bob" in sv

    def test_extract_columns_description(self, minimal_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(minimal_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert columns[0].get("description") == "Primary key"

    def test_extract_empty_columns_sheet(self):
        """A workbook with no data rows in Columns sheet returns empty list."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Columns"
        ws["A1"] = "Name"
        ws["D1"] = "Data Type"
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        assert columns == []

    def test_extract_columns_decimal_precision_scale(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(wb)
        amount_col = next(c for c in columns if c["name"] == "col_amount")
        assert amount_col.get("precision") == 10
        assert amount_col.get("scale") == 2


class TestExcelToUMFConverterExtractValidation:
    """Test _extract_validation method."""

    def test_extract_validation_rules(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        validation, review_notes = importer._extract_validation(wb)
        assert validation is not None
        expectations = validation.get("expectations", [])
        assert len(expectations) >= 2

    def test_extract_validation_rule_types(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        validation, _ = importer._extract_validation(wb)
        expectations = validation["expectations"]
        types = [e["type"] for e in expectations]
        assert "expect_column_values_to_not_be_null" in types
        assert "expect_column_values_to_be_in_set" in types

    def test_extract_validation_returns_empty_when_no_sheet(self):
        wb = openpyxl.Workbook()
        importer = ExcelToUMFConverter()
        validation, notes = importer._extract_validation(wb)
        assert validation is None
        assert notes == {}


class TestExcelToUMFConverterExtractRelationships:
    """Test _extract_relationships method."""

    def test_extract_relationships(self, rich_umf):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        importer = ExcelToUMFConverter()
        relationships = importer._extract_relationships(wb)
        assert relationships is not None
        fks = relationships.get("foreign_keys", [])
        assert len(fks) >= 1
        assert fks[0]["column"] == "col_id"
        # The exporter writes target_table/target_column, importer reads as references_table/references_column
        assert fks[0]["references_table"] == "other_table"
        assert fks[0]["references_column"] == "id"


class TestExcelToUMFConverterConvert:
    """Test full convert() method via file on disk."""

    def test_convert_file_not_found(self, tmp_path):
        importer = ExcelToUMFConverter()
        with pytest.raises(FileNotFoundError):
            importer.convert(tmp_path / "nonexistent.xlsx")

    def test_convert_invalid_workbook(self, tmp_path):
        """A workbook without required sheets should raise ValueError."""
        wb = openpyxl.Workbook()
        out = tmp_path / "bad.xlsx"
        wb.save(out)
        importer = ExcelToUMFConverter()
        with pytest.raises(ValueError, match="Invalid Excel workbook"):
            importer.convert(out)

    def test_excel_import_decimal_requires_precision_scale(self, tmp_path):
        wb = _make_manual_workbook(
            [
                {
                    "name": "col_amount",
                    "data_type": "DECIMAL",
                    "description": "Dollar amount",
                }
            ]
        )
        out = tmp_path / "decimal_missing_precision_scale.xlsx"
        wb.save(out)

        importer = ExcelToUMFConverter()
        with pytest.raises(
            ValueError,
            match=r"Columns row 2: fields precision and scale are required for DECIMAL",
        ):
            importer.convert(out)

    def test_excel_round_trip_full_convert_preserves_represented_umf_fields(
        self, tmp_path
    ):
        umf = _make_rich_umf(
            file_format=FileFormatSpec(
                delimiter="|",
                encoding="utf-8",
                header=True,
                quote_char='"',
                escape_char="\\",
                null_value="NULL",
                skip_rows=1,
                comment_char="#",
                filename_pattern=FilenamePattern(
                    regex=r"^claims_(\d{8})_(\w+)\.csv$",
                    captures={1: "as_of_date", 2: "source_system"},
                ),
            )
        )
        exporter = UMFToExcelConverter()
        workbook = exporter.convert(umf)
        out = tmp_path / "round_trip.xlsx"
        workbook.save(out)

        importer = ExcelToUMFConverter()
        round_tripped, review_notes = importer.convert(out)

        assert review_notes == {}
        validated = UMF.model_validate(round_tripped.model_dump())
        assert isinstance(validated, UMF)
        assert round_tripped.table_name == umf.table_name
        assert round_tripped.canonical_name == umf.canonical_name
        assert round_tripped.primary_key == umf.primary_key
        assert len(round_tripped.columns) == len(umf.columns)
        assert [col.name for col in round_tripped.columns] == [
            col.name for col in umf.columns
        ]

        for original, imported in zip(umf.columns, round_tripped.columns):
            assert imported.name == original.name
            assert imported.data_type == original.data_type
            assert imported.length == original.length
            assert imported.precision == original.precision
            assert imported.scale == original.scale
            assert imported.description == original.description
            assert imported.sample_values == original.sample_values
            assert imported.format == original.format
            assert imported.notes == original.notes
            assert imported.canonical_name == original.canonical_name
            assert imported.aliases == original.aliases
            assert imported.reporting_requirement == original.reporting_requirement
            if original.source is not None:
                assert imported.source == original.source
            if original.key_type is not None:
                assert imported.key_type == original.key_type
            assert (
                imported.nullable.model_dump(exclude_none=True)
                if imported.nullable
                else None
            ) == (
                original.nullable.model_dump(exclude_none=True)
                if original.nullable
                else None
            )

        def _normalize_expectations(exps: list[dict]) -> list[dict]:
            normalized = []
            for exp in exps:
                meta = dict(exp.get("meta", {}))
                meta.pop("generated_from", None)
                meta.pop("rule_index", None)
                normalized.append(
                    {
                        "type": exp.get("type"),
                        "kwargs": exp.get("kwargs", {}),
                        "meta": meta,
                    }
                )
            return sorted(normalized, key=lambda exp: json.dumps(exp, sort_keys=True))

        assert _normalize_expectations(
            round_tripped.validation_rules.expectations
        ) == _normalize_expectations(umf.validation_rules.expectations)

        assert round_tripped.relationships is not None
        assert round_tripped.relationships.foreign_keys is not None
        imported_fk = round_tripped.relationships.foreign_keys[0]
        original_rel = umf.relationships.outgoing[0]
        assert imported_fk.column == original_rel.source_column
        assert imported_fk.references_table == original_rel.target_table
        assert imported_fk.references_column == original_rel.target_column
        assert imported_fk.confidence == pytest.approx(original_rel.confidence)
        assert imported_fk.type == original_rel.type
        assert imported_fk.detection_method == "ai"

        assert round_tripped.file_format is not None
        assert round_tripped.file_format.model_dump(
            exclude_none=True
        ) == umf.file_format.model_dump(exclude_none=True)

    def test_excel_round_trip_preserves_source_and_foreign_keys(self, tmp_path):
        umf = UMF(
            version="1.0",
            table_name="orders",
            canonical_name="Orders",
            columns=[
                UMFColumn(name="order_id", data_type="INTEGER"),
                UMFColumn(name="customer_id", data_type="VARCHAR", length=10),
            ],
            source=JdbcSource(
                kind="jdbc",
                url="jdbc:sqlserver://host:1433;databaseName=db",
                dbtable="[dbo].[Orders]",
                user="svc_reader",
                password_secret_ref="scope/jdbc-password",
            ),
            relationships=Relationships(
                foreign_keys=[
                    ForeignKey(
                        column="customer_id",
                        references_table="customers",
                        references_column="customer_id",
                        confidence=1.0,
                        type="foreign_key",
                        detection_method="information_schema",
                    )
                ]
            ),
        )

        exporter = UMFToExcelConverter()
        workbook = exporter.convert(umf)
        out = tmp_path / "source_round_trip.xlsx"
        workbook.save(out)

        importer = ExcelToUMFConverter()
        round_tripped, review_notes = importer.convert(out)

        assert review_notes == {}
        assert round_tripped.source is not None
        assert round_tripped.source.model_dump(
            exclude_none=True
        ) == umf.source.model_dump(exclude_none=True)
        assert round_tripped.relationships is not None
        assert round_tripped.relationships.foreign_keys is not None
        assert [
            fk.model_dump(exclude_none=True)
            for fk in round_tripped.relationships.foreign_keys
        ] == [fk.model_dump(exclude_none=True) for fk in umf.relationships.foreign_keys]


# ---------------------------------------------------------------------------
# Round-trip tests (export then manually verify sheet content)
# ---------------------------------------------------------------------------


class TestExcelRoundTrip:
    """Test that exporting to Excel and reading back preserves data at sheet level."""

    def test_round_trip_via_file(self, rich_umf, tmp_path):
        """Export UMF -> Excel file -> reload workbook -> verify sheet data."""
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        out = tmp_path / "round_trip.xlsx"
        wb.save(out)

        wb2 = openpyxl.load_workbook(out)
        ws = wb2[ExcelConstants.SHEET_COLUMNS]

        # Verify columns are present
        col_names = []
        for row in ws.iter_rows(min_row=2, values_only=False):
            name = row[0].value
            if name:
                col_names.append(name)
        assert "col_id" in col_names
        assert "col_name" in col_names
        assert "col_amount" in col_names
        assert "col_active" in col_names
        assert "col_date" in col_names

    def test_round_trip_schema_preserved(self, rich_umf, tmp_path):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        out = tmp_path / "rt_schema.xlsx"
        wb.save(out)

        importer = ExcelToUMFConverter()
        schema = importer._extract_schema(openpyxl.load_workbook(out))
        assert schema["table_name"] == "rich_table"
        assert schema["canonical_name"] == "Rich Table"
        assert schema["version"] == "2.0"

    def test_round_trip_column_count_preserved(self, rich_umf, tmp_path):
        exporter = UMFToExcelConverter()
        wb = exporter.convert(rich_umf)
        out = tmp_path / "rt_cols.xlsx"
        wb.save(out)

        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(openpyxl.load_workbook(out))
        assert len(columns) == len(rich_umf.columns)

    def test_round_trip_nullable_preserved(self, tmp_path):
        umf = _make_minimal_umf()
        exporter = UMFToExcelConverter()
        wb = exporter.convert(umf)
        out = tmp_path / "rt_nullable.xlsx"
        wb.save(out)

        importer = ExcelToUMFConverter()
        columns = importer._extract_columns(openpyxl.load_workbook(out))
        col_name_data = next(c for c in columns if c["name"] == "col_name")
        nullable = col_name_data["nullable"]
        assert nullable["MD"] is True
        assert nullable["MP"] is False
        assert nullable["ME"] is True


# ---------------------------------------------------------------------------
# Edge cases
# ---------------------------------------------------------------------------


class TestEdgeCases:
    """Test edge cases and boundary conditions."""

    def test_single_column_umf(self):
        umf = UMF(
            version="1.0",
            table_name="one_col",
            columns=[UMFColumn(name="only_col", data_type="VARCHAR", length=50)],
        )
        converter = UMFToExcelConverter()
        wb = converter.convert(umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        assert ws["A2"].value == "only_col"
        # Row 3 should be empty
        assert ws["A3"].value is None

    def test_column_with_no_optional_fields(self):
        umf = UMF(
            version="1.0",
            table_name="bare",
            columns=[UMFColumn(name="bare_col", data_type="INTEGER")],
        )
        converter = UMFToExcelConverter()
        wb = converter.convert(umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        assert ws["A2"].value == "bare_col"
        assert ws["D2"].value == "INTEGER"
        # No nullable, no sample values
        assert ws["H2"].value is None
        assert ws["L2"].value == ""

    def test_all_data_types_export(self):
        """Verify all UMF data types can be exported to Excel."""
        cols = []
        for dtype in [
            "VARCHAR",
            "INTEGER",
            "DECIMAL",
            "DATE",
            "DATETIME",
            "BOOLEAN",
            "TEXT",
            "CHAR",
            "FLOAT",
        ]:
            extra = {}
            if dtype == "VARCHAR":
                extra = {"length": 50}
            elif dtype == "DECIMAL":
                extra = {"precision": 10, "scale": 2}
            elif dtype == "CHAR":
                extra = {"length": 1}
            cols.append(
                UMFColumn(name=f"col_{dtype.lower()}", data_type=dtype, **extra)
            )
        umf = UMF(version="1.0", table_name="all_types", columns=cols)
        converter = UMFToExcelConverter()
        wb = converter.convert(umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        exported_types = []
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value:
                exported_types.append(row[3].value)
        assert len(exported_types) == len(cols)

    def test_empty_description(self):
        umf = UMF(
            version="1.0",
            table_name="no_desc",
            columns=[UMFColumn(name="c", data_type="INTEGER")],
        )
        converter = UMFToExcelConverter()
        wb = converter.convert(umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        # Description cell should be empty string, not None
        assert ws["K2"].value in (None, "")


# ---------------------------------------------------------------------------
# Derivations sheet round-trip (UMF -> Excel -> UMF)
# ---------------------------------------------------------------------------


def _report_umf_with_derivations() -> UMF:
    """A generated report UMF with single- and multi-candidate derivations."""
    return UMF(
        version="1.0",
        table_name="member_report",
        canonical_name="Member Report",
        table_type="generated",
        description="Computed member report.",
        columns=[
            UMFColumn(
                name="id", data_type="VARCHAR", length=36, description="Primary key"
            ),
            UMFColumn(
                name="pcp_name",
                data_type="VARCHAR",
                length=200,
                description="Assigned primary care provider name",
                derivation=UMFColumnDerivation(
                    candidates=[
                        DerivationCandidate(
                            table="providers",
                            column="NAME",
                            priority=1,
                            reason="Assigned PCP.",
                            join_filter="SPECIALITY = 'GENERAL PRACTICE'",
                        ),
                        DerivationCandidate(
                            table="providers",
                            column="NAME",
                            priority=2,
                            reason="Any provider fallback.",
                        ),
                    ],
                    survivorship=Survivorship(
                        strategy="highest_priority",
                        explanation="Strategy: take the assigned PCP.\n\nFallback: any provider.",
                    ),
                ),
            ),
            UMFColumn(
                name="bmi",
                data_type="DECIMAL",
                precision=5,
                scale=2,
                description="Most recent BMI",
                derivation=UMFColumnDerivation(
                    candidates=[
                        DerivationCandidate(
                            table="observations",
                            priority=1,
                            expression="MAX(CASE WHEN DESCRIPTION='Body Mass Index' THEN VALUE END)",
                            reason="Latest BMI observation.",
                        )
                    ],
                ),
            ),
        ],
    )


class TestDerivationsRoundTrip:
    def _round_trip(self, umf: UMF, tmp_path) -> UMF:
        path = tmp_path / "report.xlsx"
        UMFToExcelConverter().convert(umf).save(path)
        rt, notes = ExcelToUMFConverter().convert(path)
        assert not {k: v for k, v in notes.items() if v}, notes
        return rt

    def test_derivations_sheet_is_written(self):
        wb = UMFToExcelConverter().convert(_report_umf_with_derivations())
        assert ExcelConstants.SHEET_DERIVATIONS in wb.sheetnames

    def test_multi_candidate_round_trip(self, tmp_path):
        rt = self._round_trip(_report_umf_with_derivations(), tmp_path)
        pcp = next(c for c in rt.columns if c.name == "pcp_name")
        assert pcp.derivation is not None
        cands = pcp.derivation.candidates or []
        assert [c.priority for c in cands] == [1, 2]
        assert cands[0].table == "providers"
        assert cands[0].column == "NAME"
        assert cands[0].join_filter == "SPECIALITY = 'GENERAL PRACTICE'"
        assert cands[1].join_filter is None
        assert pcp.derivation.survivorship is not None
        assert pcp.derivation.survivorship.strategy == "highest_priority"
        assert "Fallback" in (pcp.derivation.survivorship.explanation or "")

    def test_expression_candidate_round_trip(self, tmp_path):
        rt = self._round_trip(_report_umf_with_derivations(), tmp_path)
        bmi = next(c for c in rt.columns if c.name == "bmi")
        assert bmi.derivation is not None
        cands = bmi.derivation.candidates or []
        assert len(cands) == 1
        assert cands[0].column is None
        assert "Body Mass Index" in (cands[0].expression or "")

    def test_no_derivations_sheet_back_compat(self, tmp_path):
        # A workbook with the Derivations sheet removed imports cleanly with no
        # derivations attached (older workbooks predate the sheet).
        umf = _report_umf_with_derivations()
        path = tmp_path / "report.xlsx"
        wb = UMFToExcelConverter().convert(umf)
        del wb[ExcelConstants.SHEET_DERIVATIONS]
        wb.save(path)
        rt, notes = ExcelToUMFConverter().convert(path)
        assert all(c.derivation is None for c in rt.columns)
        assert not {k: v for k, v in notes.items() if v}

    def test_columns_without_derivation_have_none(self, tmp_path):
        rt = self._round_trip(_report_umf_with_derivations(), tmp_path)
        plain = next(c for c in rt.columns if c.name == "id")
        assert plain.derivation is None

    # -- SQL-relevant fields the SQLPlanGenerator consumes ------------------

    def test_top_level_strategy_round_trips_not_as_survivorship(self, tmp_path):
        """A primary_key-strategy column (no candidates) keeps the strategy on
        derivation.strategy, NOT survivorship.strategy."""
        umf = UMF(
            version="1.0",
            table_name="t",
            canonical_name="t",
            table_type="generated",
            columns=[
                UMFColumn(
                    name="id",
                    data_type="VARCHAR",
                    length=36,
                    derivation=UMFColumnDerivation(strategy="primary_key"),
                ),
            ],
        )
        rt = self._round_trip(umf, tmp_path)
        d = rt.columns[0].derivation
        assert d is not None
        assert d.strategy == "primary_key"
        assert d.survivorship is None

    def test_max_across_sources_strategy_round_trips(self, tmp_path):
        umf = UMF(
            version="1.0",
            table_name="t",
            canonical_name="t",
            table_type="generated",
            columns=[
                UMFColumn(
                    name="score",
                    data_type="INTEGER",
                    derivation=UMFColumnDerivation(
                        strategy="max_across_sources",
                        candidates=[
                            DerivationCandidate(table="s1", column="v", priority=1),
                            DerivationCandidate(table="s2", column="v", priority=2),
                        ],
                    ),
                ),
            ],
        )
        rt = self._round_trip(umf, tmp_path)
        assert rt.columns[0].derivation.strategy == "max_across_sources"

    def test_window_candidate_fields_round_trip(self, tmp_path):
        """row_filter / order_by / select_columns / join_via survive the trip."""
        umf = UMF(
            version="1.0",
            table_name="t",
            canonical_name="t",
            table_type="generated",
            columns=[
                UMFColumn(
                    name="a1c",
                    data_type="DECIMAL",
                    precision=5,
                    scale=2,
                    derivation=UMFColumnDerivation(
                        candidates=[
                            DerivationCandidate(
                                table="labs",
                                column="result",
                                priority=1,
                                row_filter="test_name = 'A1C'",
                                order_by=["result_date", "snapshot"],
                                select_columns=["result", "source"],
                                join_via=JoinViaSpec(
                                    lookup_table="dw",
                                    source_key="cmid",
                                    lookup_key="mid",
                                    target_key="pid",
                                ),
                                reason="latest A1C",
                            )
                        ],
                    ),
                ),
            ],
        )
        rt = self._round_trip(umf, tmp_path)
        cand = rt.columns[0].derivation.candidates[0]
        assert cand.row_filter == "test_name = 'A1C'"
        assert cand.order_by == ["result_date", "snapshot"]
        assert cand.select_columns == ["result", "source"]
        assert cand.join_via is not None
        assert cand.join_via.lookup_table == "dw"
        assert cand.join_via.target_key == "pid"

    def test_survivorship_defaults_round_trip(self, tmp_path):
        umf = UMF(
            version="1.0",
            table_name="t",
            canonical_name="t",
            table_type="generated",
            columns=[
                UMFColumn(
                    name="status",
                    data_type="VARCHAR",
                    length=20,
                    derivation=UMFColumnDerivation(
                        candidates=[
                            DerivationCandidate(table="s", column="status", priority=1)
                        ],
                        survivorship=Survivorship(
                            strategy="highest_priority",
                            explanation="Pick the highest priority source.",
                            default_value="IHA",
                            default_condition="until assessment completed",
                        ),
                    ),
                ),
            ],
        )
        rt = self._round_trip(umf, tmp_path)
        surv = rt.columns[0].derivation.survivorship
        assert surv is not None
        assert surv.default_value == "IHA"
        assert surv.default_condition == "until assessment completed"


# ---------------------------------------------------------------------------
# SQL-plan identity across an Excel round-trip
# ---------------------------------------------------------------------------


class TestDerivationSqlIdentity:
    """The strongest completeness guard: a derived table's generated gold SQL
    must be byte-identical before and after an Excel round-trip — proving no
    SQL-relevant derivation field is lost."""

    def _sql_corpus(self):
        from tablespec.schemas import generate_sql_plan

        base = UMF(
            version="1.0",
            table_name="member",
            canonical_name="member",
            table_type="ingested",
            primary_key=["member_id"],
            columns=[
                UMFColumn(name="member_id", data_type="VARCHAR", length=36),
                UMFColumn(name="bmi", data_type="DECIMAL", precision=5, scale=2),
            ],
        )
        providers = UMF(
            version="1.0",
            table_name="providers",
            canonical_name="providers",
            table_type="ingested",
            primary_key=["member_id"],
            columns=[
                UMFColumn(name="member_id", data_type="VARCHAR", length=36),
                UMFColumn(name="name", data_type="VARCHAR", length=200),
            ],
        )
        target = UMF(
            version="1.0",
            table_name="member_summary",
            canonical_name="member_summary",
            table_type="generated",
            primary_key=["member_id"],
            metadata={"base_table": "member"},
            columns=[
                UMFColumn(
                    name="member_id",
                    data_type="VARCHAR",
                    length=36,
                    derivation=UMFColumnDerivation(strategy="primary_key"),
                ),
                UMFColumn(
                    name="bmi",
                    data_type="DECIMAL",
                    precision=5,
                    scale=2,
                    derivation=UMFColumnDerivation(strategy="base_column"),
                ),
                UMFColumn(
                    name="pcp_name",
                    data_type="VARCHAR",
                    length=200,
                    derivation=UMFColumnDerivation(
                        candidates=[
                            DerivationCandidate(
                                table="providers",
                                column="name",
                                priority=1,
                                join_filter="specialty = 'GP'",
                            ),
                        ],
                        survivorship=Survivorship(
                            strategy="highest_priority",
                            explanation="GP provider name.",
                            default_value="UNKNOWN",
                        ),
                    ),
                ),
            ],
        )
        related = {"member": base, "providers": providers}
        return generate_sql_plan, target, related

    def test_gold_sql_identical_after_round_trip(self, tmp_path):
        generate_sql_plan, target, related = self._sql_corpus()
        before = generate_sql_plan(target, related, mode="cte")

        path = tmp_path / "member_summary.xlsx"
        UMFToExcelConverter().convert(target).save(path)
        rt, notes = ExcelToUMFConverter().convert(path)
        assert not {k: v for k, v in notes.items() if v}, notes

        after = generate_sql_plan(rt, related, mode="cte")
        assert after == before


# ---------------------------------------------------------------------------
# Data-validation limits (Excel rejects inline list formulas over 255 chars)
# ---------------------------------------------------------------------------


class TestDataValidationLimits:
    """Excel silently strips (and reports as 'recovered') any data-validation
    whose inline string-literal formula exceeds 255 characters. Long option
    lists (e.g. domain types) must reference a range instead."""

    def test_no_formula1_exceeds_excel_limit(self):
        # A column with a domain_type triggers the domain-type dropdown, the
        # longest option list the exporter emits.
        umf = UMF(
            version="1.0",
            table_name="dv_table",
            canonical_name="dv_table",
            columns=[
                UMFColumn(
                    name="state",
                    data_type="VARCHAR",
                    length=2,
                    domain_type="us_state_code",
                ),
            ],
        )
        wb = UMFToExcelConverter().convert(umf)
        offenders: list[tuple[str, int]] = []
        for ws in wb.worksheets:
            for dv in ws.data_validations.dataValidation:
                f1 = dv.formula1 or ""
                # Inline list = a quoted string literal; range refs are exempt.
                if f1.startswith('"') and len(f1) > 255:
                    offenders.append((ws.title, len(f1)))
        assert not offenders, f"data-validation formula1 over 255 chars: {offenders}"

    def test_domain_type_dropdown_references_instructions_range(self):
        umf = UMF(
            version="1.0",
            table_name="dv_table",
            canonical_name="dv_table",
            columns=[
                UMFColumn(name="c", data_type="VARCHAR", domain_type="email"),
            ],
        )
        wb = UMFToExcelConverter().convert(umf)
        ws = wb[ExcelConstants.SHEET_COLUMNS]
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        assert any(
            (f or "").startswith(f"'{ExcelConstants.SHEET_INSTRUCTIONS}'!$P$")
            for f in formulas
        ), formulas


class TestPipelineBlocksRoundTrip:
    """ingestion + typed metadata (source_tables, output_config) round-trip."""

    def _round_trip(self, umf, tmp_path):
        from tablespec.excel_converter import ExcelToUMFConverter, UMFToExcelConverter

        path = tmp_path / "rt.xlsx"
        UMFToExcelConverter().convert(umf).save(path)
        back, _ = ExcelToUMFConverter().convert(path)
        return back

    def test_ingestion_and_source_path_round_trip(self, tmp_path):
        from tablespec.models.umf import UMF

        umf = UMF.model_validate(
            {
                "version": "1.0",
                "table_name": "orders",
                "canonical_name": "orders",
                "primary_key": ["order_id"],
                "ingestion": {"mode": "incremental", "order_by": ["_load_ts"]},
                "source": {
                    "kind": "delimited",
                    "delimiter": "|",
                    "header": True,
                    "path": "/data/orders_*.csv",
                    "filename_pattern": {
                        "regex": r"orders_(\d{8})\.csv",
                        "captures": {1: "file_dt"},
                    },
                },
                "columns": [
                    {
                        "name": "order_id",
                        "data_type": "VARCHAR",
                        "nullable": {"default": False},
                    }
                ],
            }
        )
        back = self._round_trip(umf, tmp_path)
        assert back.primary_key == ["order_id"]
        assert back.ingestion is not None
        assert back.ingestion.mode == "incremental"
        assert back.ingestion.order_by == ["_load_ts"]
        assert back.source is not None and back.source.path == "/data/orders_*.csv"
        assert back.source.filename_pattern.captures == {1: "file_dt"}

    def test_structured_metadata_round_trips_typed(self, tmp_path):
        from tablespec.models.umf import UMF

        umf = UMF.model_validate(
            {
                "version": "1.0",
                "table_name": "report",
                "canonical_name": "report",
                "table_type": "generated",
                "metadata": {
                    "base_table_strategy": "union_sources",
                    "source_tables": ["a", "b"],
                    "output_config": {
                        "include_footer": True,
                        "line_terminator": "CRLF",
                        "file_naming_example": "report_YYYYMMDD.csv",
                    },
                },
                "columns": [{"name": "k", "data_type": "VARCHAR"}],
            }
        )
        back = self._round_trip(umf, tmp_path)
        assert back.metadata is not None
        assert back.metadata.base_table_strategy == "union_sources"
        assert back.metadata.source_tables == ["a", "b"]
        config = back.metadata.output_config
        assert config is not None and config.include_footer is True
        assert config.line_terminator == "CRLF"
        assert config.file_naming_example == "report_YYYYMMDD.csv"
